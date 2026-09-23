"""Archive every CNB Monetary Policy Report macro-indicator table and parse all of its rows.

The existing `data/cnb_mpr_cpi_quarterly.csv` keeps only the headline CPI row and does not keep
the source spreadsheets. This tool downloads each report's spreadsheet once, stores the exact
bytes with a hash, and writes every indicator row (prices by component, wages, rates, exchange
rates and the external assumptions) in long form. A cell is a CNB forecast when it is printed
in bold, as the table's own footnote says.

These tables are CNB forecasts. They are evaluation references and inputs to the separate,
clearly labelled conditioned lane only. Nothing here may enter the independent forecast.

    python -m tools.cnb_tracker.fetch_tables --output data/cnb_mpr_tables_YYYYMMDD
"""
import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import time
import urllib.request

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]; sys.path.insert(0, str(ROOT))
REGISTRY = ROOT / 'data/cnb_mpr_cpi_quarterly.csv'
QUARTER = {'Q1': 1, 'Q2': 2, 'Q3': 3, 'Q4': 4, 'QI': 1, 'QII': 2, 'QIII': 3, 'QIV': 4, 'I': 1, 'II': 2, 'III': 3, 'IV': 4}
KEYS = [  # (key, label must start with, label must also contain)
    ('cpi', 'consumer price index', 'y-o-y'), ('administered', 'administered prices', ''), ('food_alc_tobacco', 'food prices', ''),
    ('core', 'core inflation', ''), ('fuel', 'fuel prices', ''), ('mp_relevant', 'monetary policy-relevant inflation', ''),
    ('ppi', 'industrial producer prices', ''), ('agri_prices', 'agricultural prices', ''), ('gdp_deflator', 'gdp deflator', ''),
    ('gdp_yy', 'gdp', 'y-o-y'), ('output_gap', 'output gap', ''), ('wage_nominal', 'average monthly wage', 'nominal terms'),
    ('wage_market', 'average monthly wage in market sectors', ''), ('ulc', 'unit labour costs', ''), ('productivity', 'aggregate labour productivity', ''),
    ('unemployment_ilo', 'ilo general unemployment rate', ''), ('repo_2w', '2w repo rate', ''), ('pribor_3m', '3m pribor', ''),
    ('czk_usd', 'czk/usd', ''), ('czk_eur', 'czk/eur', ''), ('euribor_3m', '3m euribor', ''), ('usd_eur', 'usd/eur', ''),
    ('foreign_gdp_yy', 'foreign gdp', 'y-o-y'), ('foreign_hicp', 'foreign hicp', ''), ('foreign_ppi', 'foreign ppi', ''), ('brent', 'brent crude oil', '')]


def indicator_key(label):
    text = re.sub(r'\s+', ' ', str(label)).strip().lower()
    hits = [key for key, start, also in KEYS if text.startswith(start) and also in text]
    if 'wage_market' in hits:
        return 'wage_market'
    if text.startswith('average monthly wage') and ('non-market' in text or 'real terms' in text):
        return None
    return hits[0] if hits else None


def parse(content):
    """All numeric cells of the English sheet: label, frequency, period, value, bold (= CNB forecast)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb['English version'] if 'English version' in wb.sheetnames else wb.worksheets[0]
    grid = [list(row) for row in ws.iter_rows()]
    header = next((i for i, row in enumerate(grid) if sum(isinstance(c.value, (int, float)) and 2000 <= c.value <= 2040 for c in row) >= 5
                   and any(str(c.value).strip() in QUARTER for c in row)), None)
    if header is None:
        raise ValueError('No header row with years and quarters')
    years_above = grid[header - 1]; columns = {}; current = None
    for j, cell in enumerate(grid[header]):
        above = years_above[j].value if j < len(years_above) else None
        if isinstance(above, (int, float)) and 2000 <= above <= 2040:
            current = int(above)
        value = cell.value
        if isinstance(value, (int, float)) and 2000 <= value <= 2040:
            columns[j] = ('A', str(int(value)))
        elif isinstance(value, str) and value.strip() in QUARTER and current is not None:
            columns[j] = ('Q', f'{current}Q{QUARTER[value.strip()]}')
    if sum(f == 'Q' for f, _ in columns.values()) < 8 or len(set(columns.values())) != len(columns):
        raise ValueError('Unexpected quarterly layout')
    rows = []; section = None
    for row in grid[header + 1:]:
        label = next((str(c.value).strip() for c in row[:3] if isinstance(c.value, str) and c.value.strip()), None)
        if label is None:
            continue
        cells = [(j, row[j]) for j in columns if j < len(row) and isinstance(row[j].value, (int, float)) and not isinstance(row[j].value, bool)]
        if not cells:
            section = label if label.isupper() else section
            continue
        for j, cell in cells:
            frequency, period = columns[j]
            rows.append(dict(section=section, row_label=re.sub(r'\s+', ' ', label), indicator=indicator_key(label), frequency=frequency, period=period,
                             value=float(cell.value), is_forecast=bool(cell.font is not None and cell.font.b)))
    return pd.DataFrame(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__); ap.add_argument('--output', type=Path, required=True); args = ap.parse_args()
    out = args.output; (out / 'raw').mkdir(parents=True, exist_ok=False)
    registry = pd.read_csv(REGISTRY); reports = registry.drop_duplicates('report_date')[['report_date', 'cutoff_date', 'season', 'vintage_year', 'xlsx_url']]
    frames = []; sources = []
    for r in reports.itertuples():
        request = urllib.request.Request(r.xlsx_url, headers={'User-Agent': 'Mozilla/5.0 (research archive of published CNB forecast tables)'})
        content = urllib.request.urlopen(request, timeout=90).read(); name = f'{r.report_date}_{r.season}_{r.vintage_year}.xlsx'
        (out / 'raw' / name).write_bytes(content)
        sources.append(dict(file='raw/' + name, url=r.xlsx_url, bytes=len(content), sha256=hashlib.sha256(content).hexdigest(),
                            retrieved_at_utc=datetime.now(timezone.utc).isoformat()))
        table = parse(content); table.insert(0, 'report_date', r.report_date); table.insert(1, 'cutoff_date', r.cutoff_date)
        table.insert(2, 'season', r.season); table.insert(3, 'vintage_year', r.vintage_year); frames.append(table)
        print(r.report_date, r.season, r.vintage_year, len(table), 'cells;', table.indicator.nunique(), 'keyed indicators', flush=True); time.sleep(.4)
    long = pd.concat(frames, ignore_index=True)
    # The parser must reproduce the frozen headline benchmark exactly before anything else is trusted.
    mine = long[long.indicator.eq('cpi') & long.frequency.eq('Q')].set_index(['report_date', 'period']).value
    frozen = registry.set_index(['report_date', 'quarter']).value
    if not mine.index.sort_values().equals(frozen.index.sort_values()):
        raise AssertionError('Parsed CPI quarters differ from data/cnb_mpr_cpi_quarterly.csv')
    np.testing.assert_allclose(mine.reindex(frozen.index).to_numpy(), frozen.to_numpy(), atol=1e-12, rtol=0)
    long.to_csv(out / 'cnb_mpr_indicators_long.csv', index=False)
    (out / 'source_manifest.json').write_text(json.dumps(dict(
        created_at_utc=datetime.now(timezone.utc).isoformat(), registry='data/cnb_mpr_cpi_quarterly.csv',
        registry_sha256=hashlib.sha256(REGISTRY.read_bytes()).hexdigest(), reports=len(reports), sources=sources,
        outputs={'cnb_mpr_indicators_long.csv': hashlib.sha256((out / 'cnb_mpr_indicators_long.csv').read_bytes()).hexdigest()},
        forecast_flag='bold cell, per the table footnote "data in bold = CNB forecast"',
        check='parsed quarterly CPI equals data/cnb_mpr_cpi_quarterly.csv exactly for every report',
        use='CNB forecasts: evaluation references and the separate conditioned lane only; never an input to the independent forecast'), indent=2), encoding='utf-8')
    print('Archived', len(reports), 'reports;', len(long), 'cells ->', out, flush=True)


if __name__ == '__main__':
    main()
