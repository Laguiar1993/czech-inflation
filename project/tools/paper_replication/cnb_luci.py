"""CNB LUCI for Table A6 rows 15-16: a tidy quarterly file from the ARAD download.

The raw files in ``data/paper_replication/official_cnb_luci_20260912/raw`` were
downloaded on 12 September 2026. ``sources.json`` in that folder records the URL,
request body, HTTP status and SHA-256 of every file. This step:

* checks the ARAD JSON and the Summer 2026 chart-data workbook against those
  hashes;
* reads the Monetary Policy Report Summer 2026 baseline snapshot (ARAD id 95) of
  the LUCI total (``MLUCLUTXXINDQ``) and its five contributions, including wages
  and costs (``MLUCLUWXXSTDQ``). Every series must equal the report workbook
  column of the same chart to ARAD's two-decimal rounding;
* drops the CNB forecast, since a report published in quarter P observes quarters
  up to P-2;
* writes ``luci_quarterly.csv`` with a declared availability date per quarter.

Availability rule (declared; these are not recorded publication times): quarter q
is first published in the CNB report of quarter q+2, on the report date listed in
``data/cnb_mpr_cpi_quarterly.csv``. For quarters before those dates the rule is
the 15th of the report month (February, May, August or November) of q+2. All
values are the current (Summer 2026) vintage, as for every other A6 input.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT / "data" / "paper_replication" / "official_cnb_luci_20260912"
REPORT_DATES = ROOT / "data" / "cnb_mpr_cpi_quarterly.csv"
ARAD_FILE = "raw/arad_indicators-data-by-codes_LUCI_6codes.json"
WORKBOOK = "raw/zomp_2026_leto_podkladova_data.xlsx"
SNAPSHOT_ID = 95
REPORT = ("summer", 2026)
SERIES = {  # tidy column: (ARAD code, report chart-data label)
    "luci_total": ("MLUCLUTXXINDQ", "II_9_LUCI"),
    "luci_wages_labour_costs": ("MLUCLUWXXSTDQ", "II_9_Wages_and_costs"),
    "luci_employment": ("MLUCLUEXXSTDQ", "II_9_Employment"),
    "luci_unemployment": ("MLUCLUUXXSTDQ", "II_9_Unemployment"),
    "luci_demand_for_labour": ("MLUCLUDXXSTDQ", "II_9_Demand_for_labour"),
    "luci_other": ("MLUCLURXXSTDQ", "II_9_Other"),
}
ROUNDING_TOL = 0.0051
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4}
RULES = {
    "values": "ARAD snapshot 95 (Monetary Policy Report Summer 2026 baseline), two decimals, current vintage",
    "observed_quarters": "a report published in quarter P observes quarters up to P-2; later quarters are CNB forecasts and are dropped",
    "available_from_assumed": ("report date of the CNB report published in quarter q+2 (data/cnb_mpr_cpi_quarterly.csv); "
                               "before the listed dates, the 15th of the report month (Feb/May/Aug/Nov) of q+2"),
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_hashes(folder: Path, names: list[str]) -> dict[str, str]:
    recorded = {f["file"]: f["sha256"] for f in json.loads((folder / "sources.json").read_text(encoding="utf-8"))["files"]}
    verified = {}
    for name in names:
        actual = sha256_file(folder / name)
        if recorded.get(name) != actual:
            raise ValueError(f"{name}: SHA-256 {actual} does not match sources.json ({recorded.get(name)})")
        verified[name] = actual
    return verified


def roman_quarter(label) -> pd.Period | None:
    if not isinstance(label, str) or "/" not in label:
        return None
    numeral, year = label.strip().split("/")
    return pd.Period(year=int(year), quarter=ROMAN[numeral], freq="Q")


def read_arad(path: Path, snapshot_id: int = SNAPSHOT_ID) -> pd.DataFrame:
    """Quarterly values of every indicator in an ARAD ``indicators-data-by-codes`` response, for one snapshot."""
    payload = json.loads(Path(path).read_text(encoding="utf-8"))["data"][0]
    columns = {}
    for indicator in payload["indicators"]:
        chosen = [s for s in indicator["snapshots_data"] if s["id"] == snapshot_id]
        if len(chosen) != 1:
            raise ValueError(f"{indicator['code']}: snapshot {snapshot_id} found {len(chosen)} times")
        points = chosen[0]["data"]
        stamps = pd.to_datetime([p[0] for p in points], unit="ms")
        quarters = stamps.to_period("Q")
        if not (stamps == quarters.to_timestamp(how="end").normalize()).all():
            raise ValueError(f"{indicator['code']}: dates are not quarter ends")
        values = [np.nan if p[1] is None else p[1] for p in points]
        columns[indicator["code"]] = pd.Series(values, index=quarters, dtype=float)
    return pd.DataFrame(columns).sort_index()


def read_workbook(path: Path, sheet: str = "QUARTERLY_DATA") -> pd.DataFrame:
    """Report chart-data columns located by label (either label row), indexed by quarter."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[sheet].iter_rows(values_only=True))
    workbook.close()
    where: dict[str, int] = {}
    for row in rows[:2]:
        for i, label in enumerate(row):
            if label is not None:
                where.setdefault(str(label).strip(), i)
    body = [r for r in rows[2:] if r and roman_quarter(r[0]) is not None]
    index = pd.PeriodIndex([roman_quarter(r[0]) for r in body], freq="Q")
    frame = {}
    for name, (_, label) in SERIES.items():
        if label not in where:
            raise ValueError(f"{Path(path).name}: label {label} not found")
        i = where[label]
        frame[name] = pd.to_numeric(pd.Series([r[i] if i < len(r) else None for r in body], index=index), errors="coerce")
    return pd.DataFrame(frame)


def report_date(dates_csv: Path, season: str, year: int) -> pd.Timestamp:
    frame = pd.read_csv(dates_csv)[["report_date", "season", "vintage_year"]].drop_duplicates()
    hit = frame[frame.season.eq(season) & frame.vintage_year.eq(year)]
    if hit.report_date.nunique() != 1:
        raise ValueError(f"{dates_csv}: expected one report date for {season} {year}, found {hit.report_date.nunique()}")
    return pd.Timestamp(hit.report_date.iloc[0])


def last_observed_quarter(published: pd.Timestamp) -> pd.Period:
    return pd.Period(published, freq="Q") - 2


def available_from(quarters: pd.PeriodIndex, report_dates) -> list[pd.Timestamp]:
    first_in_quarter: dict[pd.Period, pd.Timestamp] = {}
    for date in sorted(pd.to_datetime(pd.Series(list(report_dates))).dropna()):
        first_in_quarter.setdefault(pd.Period(date, freq="Q"), pd.Timestamp(date))
    out = []
    for quarter in quarters:
        publication = quarter + 2
        date = first_in_quarter.get(publication)
        if date is None:
            date = pd.Timestamp(year=publication.year, month=publication.start_time.month + 1, day=15)
        out.append(date)
    return out


def build(folder: Path = FOLDER, dates_csv: Path = REPORT_DATES, replace: bool = False) -> dict:
    folder = Path(folder)
    output = folder / "luci_quarterly.csv"
    if output.exists() and not replace:
        raise FileExistsError(output)
    hashes = verify_hashes(folder, [ARAD_FILE, WORKBOOK])
    arad = read_arad(folder / ARAD_FILE)
    codes = {code: name for name, (code, _) in SERIES.items()}
    missing = sorted(set(codes) - set(arad.columns))
    if missing:
        raise ValueError(f"ARAD response lacks {missing}")
    tidy = arad[list(codes)].rename(columns=codes)
    chart = read_workbook(folder / WORKBOOK)
    common = tidy.index.intersection(chart.index)
    differences = {name: float((tidy.loc[common, name] - chart.loc[common, name]).abs().max()) for name in SERIES}
    beyond = {k: v for k, v in differences.items() if not v <= ROUNDING_TOL}
    if beyond:
        raise ValueError(f"ARAD differs from the report workbook beyond rounding: {beyond}")
    parts = [name for name in SERIES if name != "luci_total"]
    component_gap = float((tidy[parts].sum(axis=1) - tidy.luci_total).abs().max())
    published = report_date(dates_csv, *REPORT)
    last = last_observed_quarter(published)
    observed = tidy.loc[:last].dropna(how="all").copy()
    dropped = tidy.index[tidy.index > last]
    report_dates = pd.read_csv(dates_csv).report_date.unique()
    observed.insert(0, "period", observed.index.astype(str))
    observed["available_from_assumed"] = [d.date().isoformat() for d in available_from(observed.index, report_dates)]
    observed["source"] = f"CNB ARAD snapshot {SNAPSHOT_ID} (Monetary Policy Report Summer 2026 baseline)"
    observed.to_csv(output, index=False)
    manifest = {
        "tool": "cnb_luci", "created_at": datetime.now(timezone.utc).isoformat(),
        "inputs_verified_against_sources_json": hashes,
        "report_dates_file": {"path": str(dates_csv), "sha256": sha256_file(dates_csv)},
        "arad_codes": {name: code for name, (code, _) in SERIES.items()},
        "snapshot_id": SNAPSHOT_ID, "report": f"{REPORT[0]} {REPORT[1]}", "report_date": published.date().isoformat(),
        "last_observed_quarter": str(last),
        "forecast_quarters_dropped": {"n": len(dropped), "first": str(dropped.min()) if len(dropped) else None,
                                      "last": str(dropped.max()) if len(dropped) else None},
        "rules": RULES,
        "checks": {"max_abs_difference_vs_report_workbook": differences,
                   "workbook_quarters_compared": {"n": len(common), "first": str(common.min()), "last": str(common.max())},
                   "rounding_tolerance": ROUNDING_TOL,
                   "max_abs_gap_components_sum_vs_total": component_gap},
        "rows": len(observed), "first_period": str(observed.index.min()), "last_period": str(observed.index.max()),
        "output": {"file": output.name, "sha256": sha256_file(output)},
        "tool_sha256": sha256_file(Path(__file__)),
    }
    (folder / "luci_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--folder", type=Path, default=FOLDER)
    parser.add_argument("--replace", action="store_true", help="rewrite an existing luci_quarterly.csv")
    args = parser.parse_args(argv)
    manifest = build(args.folder, replace=args.replace)
    print(json.dumps({k: manifest[k] for k in ("rows", "first_period", "last_period", "last_observed_quarter",
                                                "forecast_quarters_dropped", "checks")}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
