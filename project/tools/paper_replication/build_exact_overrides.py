"""Exact Table A6 inputs from CNB ARAD and CZSO, for rows that used stand-ins or truncated histories.

Matched on 12 September 2026. The paper lists its sources as Eurostat, ARAD CNB, CZSO and
Refinitiv, and ARAD carries series whose labels match Table A6 rows. The downloads, with
request bodies, HTTP headers and SHA-256, are in
``data/paper_replication/official_cnb_arad_20260912/raw`` and
``official_czso_20260912_permits/raw``.

Rows written (value kinds follow ``exact_inputs_long.csv``):

* 13: CZSO Table 6 building permits granted, monthly count
* 17: ARAD ``MLULNULXXADJYOYPECQ`` nominal unit labour costs, y/y %, seasonally
  adjusted, quarterly. Table A6 leaves row 17 untransformed, which fits a
  stationary rate rather than the level index used before. ARAD has no values
  for 2001, so the series starts in 2002Q1, before the paper sample.
* 25: ARAD ``SVEVZM4`` balance of foreign trade (FOB/FOB)
* 26: ARAD ``SIMCSUM2005PM01`` (2000-01 to 2017-12) and ``SIMCSUM2015PM01``
  (2018 on), import prices, previous month = 100
* 43: ARAD ``MOPAAPPXXNAJYOYPECM`` agricultural producer prices, y/y %
* 52: CZSO CEN02032 agriculture including fish, previous month = 100, from
  2010-01, extended back with ARAD ``MOPAAPPXXNAJYOYPECM`` y/y rates
* 54: ARAD ``SFTP04M2106`` PRIBOR 3M, end of month
* 58 and 59: ARAD ``SUCM100311XXX101101`` / ``SUCM102211XXX101101`` client loans
* 60 to 63: ARAD ``MEDACOMOILXXUSBVALM``, ``MEDACOMGASXXL18INDM``,
  ``MEDACOMMETXXL18INDM`` and ``MEDACOMFODXXL18INDM``

Report-database series (``MEDACOM*``, ``MOPA*``, ``MLULNUL*``) come in the Summer 2026 report
snapshot, which runs to the end of the forecast horizon. Only observed periods are kept: to
2026-06, where ARAD Brent stops equalling Bloomberg, and to 2026Q1 for quarterly series. The
availability rules (days after the period end) are those of the rows being replaced.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "paper_replication"
ARAD_RAW = DATA / "official_cnb_arad_20260912" / "raw"
ARAD_FILES = {"candidates": "arad_a6_candidates_14codes", "prices": "arad_price_indices"}
PERMITS = DATA / "official_czso_20260912_permits" / "raw" / "bvzcr090726_06.xlsx"
INPUTS = DATA / "a6_inputs_20260912"
DEFAULT_OUTPUT = DATA / "a6_exact_overrides_20260912"
LAST_MONTH = pd.Period("2026-06", freq="M")
LAST_QUARTER = pd.Period("2026Q1", freq="Q")
PANEL_START = pd.Period("2002-05", freq="M")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

# (row, download, ARAD codes, value kind, frequency, report database, availability lag in days, description)
SPEC = [
    (17, "candidates", ("MLULNULXXADJYOYPECQ",), "level", "Q", True, 75, "nominal unit labour costs, y/y %, SA"),
    (25, "candidates", ("SVEVZM4",), "level", "M", False, 40, "balance of foreign trade (FOB/FOB)"),
    (26, "prices", ("SIMCSUM2005PM01", "SIMCSUM2015PM01"), "mm_index_previous_month_100", "M", False, 45, "import prices, previous month = 100"),
    (43, "candidates", ("MOPAAPPXXNAJYOYPECM",), "level", "M", True, 26, "agricultural producer prices, y/y %"),
    (54, "candidates", ("SFTP04M2106",), "level", "M", False, 1, "PRIBOR 3M, end of month"),
    (58, "candidates", ("SUCM100311XXX101101",), "level", "M", False, 31, "client loans, non-financial corporations"),
    (59, "candidates", ("SUCM102211XXX101101",), "level", "M", False, 31, "client loans, households"),
    (60, "candidates", ("MEDACOMOILXXUSBVALM",), "level", "M", True, 1, "Brent crude oil, USD/barrel"),
    (61, "candidates", ("MEDACOMGASXXL18INDM",), "level", "M", True, 1, "average natural gas price in Europe, 2018 = 100"),
    (62, "candidates", ("MEDACOMMETXXL18INDM",), "level", "M", True, 1, "industrial metals price index, 2018 = 100"),
    (63, "candidates", ("MEDACOMFODXXL18INDM",), "level", "M", True, 1, "food commodity price index, 2018 = 100"),
]
POLICIES = {17: "SA year-on-year rate published by the CNB (ARAD); do not re-adjust"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_download(stem: str, folder: Path = ARAD_RAW) -> dict:
    raw, request = folder / f"{stem}.json", json.loads((folder / f"{stem}.request.json").read_text(encoding="utf-8"))
    actual = sha256_file(raw)
    if request["sha256"] != actual:
        raise ValueError(f"{raw.name}: SHA-256 {actual} differs from the recorded {request['sha256']}")
    return {"file": str(raw.relative_to(ROOT)), "sha256": actual, "url": request["url"], "body": request["body"],
            "retrieved_utc": request["retrieved_utc"]}


def arad_series(path: Path, code: str, freq: str = "M") -> tuple[pd.Series, object]:
    """Values of one indicator in an ARAD ``indicators-data-by-codes`` response (first snapshot) and its snapshot id."""
    for indicator in json.loads(Path(path).read_text(encoding="utf-8"))["data"][0]["indicators"]:
        if indicator["code"] == code:
            snapshot = indicator["snapshots_data"][0]
            points = [p for p in snapshot["data"] if p[1] is not None]
            stamps = pd.to_datetime([p[0] for p in points], unit="ms")
            series = pd.Series([p[1] for p in points], index=stamps.to_period(freq), dtype=float)
            return series[~series.index.duplicated(keep="last")].sort_index(), snapshot.get("id")
    raise KeyError(code)


def czso_permits(path: Path = PERMITS, sheet: str = "od roku 2020") -> pd.Series:
    """Monthly number of building permits granted from CZSO Table 6 (rows labelled 'YYYY - Leden / January', then month names)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    values, year = {}, None
    for row in workbook[sheet].iter_rows(values_only=True):
        if not row or not isinstance(row[0], str) or "/" not in row[0]:
            continue
        match = re.match(r"\s*(\d{4})\s*-", row[0])
        if match:
            year = int(match.group(1))
        english = row[0].split("/")[-1].strip()
        if year is not None and english in MONTHS and row[1] is not None:
            values[pd.Period(year=year, month=MONTHS.index(english) + 1, freq="M")] = float(row[1])
    workbook.close()
    return pd.Series(values, dtype=float).sort_index()


def chain_previous_month(parts: list[pd.Series]) -> pd.Series:
    """Previous-month indices from successive base years; a later base wins where two overlap."""
    out = parts[0]
    for part in parts[1:]:
        out = part.combine_first(out)
    return out.sort_index()


def contiguous_tail(series: pd.Series) -> tuple[pd.Series, list[str]]:
    """The part of a series after its last internal gap, and the missing periods."""
    full = pd.period_range(series.index.min(), series.index.max(), freq=series.index.freq)
    gaps = full.difference(series.index)
    return (series.loc[gaps.max() + 1:] if len(gaps) else series), [str(p) for p in gaps]


def backcast_previous_month(mm_index: pd.Series, yoy: pd.Series, start: pd.Period) -> pd.Series:
    """Extend a previous-month index (= 100) back to ``start`` with year-on-year rates.

    Levels are anchored on the observed chain, and L[s] = L[s + 12] / (1 + yoy[s + 12] / 100)
    walking back to ``start - 1``.
    """
    first = mm_index.index.min()
    level = (mm_index / 100.0).cumprod()
    level[first - 1] = 1.0
    for s in pd.period_range(start - 1, first - 2, freq="M")[::-1]:
        level[s] = level[s + 12] / (1.0 + yoy[s + 12] / 100.0)
    level = level.sort_index()
    rebuilt = (level / level.shift(1) * 100.0).loc[start:first - 1]
    return pd.concat([rebuilt, mm_index]).sort_index()


def backcast_check(mm_index: pd.Series, yoy: pd.Series, anchor: str = "2015-01", test=("2010-02", "2014-12")) -> dict:
    """Rebuild the observed months before ``anchor`` from y/y rates and compare the monthly changes."""
    rebuilt = backcast_previous_month(mm_index.loc[anchor:], yoy, mm_index.index.min())
    actual, estimate = np.log(mm_index.loc[test[0]:test[1]] / 100.0), np.log(rebuilt.loc[test[0]:test[1]] / 100.0)
    return {"months": int(len(actual)), "window": list(test), "mae_pp": float((actual - estimate).abs().mean() * 100),
            "max_abs_pp": float((actual - estimate).abs().max() * 100), "corr": float(actual.corr(estimate))}


def rows_frame(number: int, series: pd.Series, source_id: str, value_kind: str, frequency: str, lag_days: int,
               source_file: str, policy: str | None = None) -> pd.DataFrame:
    series = series.dropna()
    ends = series.index.to_timestamp(how="end").normalize()
    return pd.DataFrame({
        "a6_number": number, "component": "", "source_id": source_id, "source_file": source_file,
        "period": series.index.astype(str), "value": series.to_numpy(float), "value_kind": value_kind,
        "availability_rule": f"t_plus_{lag_days}d", "available_from_assumed": (ends + pd.Timedelta(days=lag_days)).strftime("%Y-%m-%d"),
        "vintage": "current, downloaded 2026-09-12", "frequency": frequency, "seasonal_adjustment": policy})


def _long(frame: pd.DataFrame, freq: str = "M", **match) -> pd.Series:
    rows = frame
    for column, value in match.items():
        rows = rows[rows[column].eq(value)]
    return pd.Series(rows.value.to_numpy(float), index=pd.PeriodIndex(rows.period, freq=freq)).groupby(level=0).last().sort_index()


def _overlap(a: pd.Series, b: pd.Series, log_changes: bool = False) -> dict:
    joined = pd.concat([a, b], axis=1, keys=["new", "old"]).dropna()
    if joined.empty:
        return {"months": 0}
    out = {"months": int(len(joined)), "first": str(joined.index.min()), "last": str(joined.index.max())}
    if log_changes:
        changes = np.log(joined).diff().dropna()
        out["log_change_corr"] = float(changes.new.corr(changes.old))
    else:
        out.update(mae=float((joined.new - joined.old).abs().mean()), max_abs=float((joined.new - joined.old).abs().max()),
                   corr=float(joined.new.corr(joined.old)), median_ratio=float((joined.new / joined.old).median()))
    return out


def build(output: Path = DEFAULT_OUTPUT, replace: bool = False) -> dict:
    output = Path(output)
    if output.exists() and not replace:
        raise FileExistsError(output)
    output.mkdir(parents=True, exist_ok=True)
    downloads = {key: verify_download(stem) for key, stem in ARAD_FILES.items()}
    frames, sources, snapshots = [], {}, {}
    for number, key, codes, kind, freq, report_database, lag, description in SPEC:
        parts = []
        for code in codes:
            series, snapshot = arad_series(ARAD_RAW / f"{ARAD_FILES[key]}.json", code, freq)
            parts.append(series)
            snapshots[code] = snapshot
        series = chain_previous_month(parts) if len(parts) > 1 else parts[0]
        if report_database:
            series = series.loc[:LAST_QUARTER if freq == "Q" else LAST_MONTH]
        gaps = []
        if freq == "Q":   # Chow-Lin needs an unbroken quarterly series
            series, gaps = contiguous_tail(series)
        source_id = "cnb_arad:" + "+".join(codes)
        frames.append(rows_frame(number, series, source_id, kind, freq, lag, downloads[key]["file"], POLICIES.get(number)))
        sources[number] = {"source_id": source_id, "description": description, "first": str(series.index.min()),
                           "last": str(series.index.max()), "observations": int(series.notna().sum()), "lag_days": lag,
                           "missing_source_periods_dropped_with_earlier_history": gaps}
    permits = czso_permits()
    frames.append(rows_frame(13, permits, "czso:bvzcr_tab6:building_permits_granted", "level", "M", 41, str(PERMITS.relative_to(ROOT))))
    sources[13] = {"source_id": "czso:bvzcr_tab6:building_permits_granted", "description": "building permits granted, monthly count",
                   "first": str(permits.index.min()), "last": str(permits.index.max()), "observations": int(permits.notna().sum()), "lag_days": 41}
    exact = pd.read_csv(INPUTS / "exact_inputs_long.csv", dtype={"period": str, "component": str})
    candidates = pd.read_csv(INPUTS / "candidate_inputs_long.csv", dtype={"period": str})
    mm52 = _long(exact, a6_number=52)
    yoy_total, _ = arad_series(ARAD_RAW / f"{ARAD_FILES['candidates']}.json", "MOPAAPPXXNAJYOYPECM")
    extended = backcast_previous_month(mm52, yoy_total, PANEL_START)
    source_52 = "czso:CEN02032:0:IM; before 2010-01 rebuilt from cnb_arad:MOPAAPPXXNAJYOYPECM"
    frames.append(rows_frame(52, extended, source_52, "mm_index_previous_month_100", "M", 26, "data/paper_replication/a6_inputs_20260912/exact_inputs_long.csv"))
    sources[52] = {"source_id": source_52, "description": "agricultural producer prices incl. fish, previous month = 100",
                   "first": str(extended.index.min()), "last": str(extended.index.max()), "observations": int(extended.notna().sum()), "lag_days": 26}
    frame = pd.concat(frames, ignore_index=True).sort_values(["a6_number", "period"], kind="mergesort")
    frame.to_csv(output / "overrides_long.csv", index=False)

    new = {n: _long(frame, "Q" if n == 17 else "M", a6_number=n) for n in frame.a6_number.unique()}
    checks = {
        "13_vs_bloomberg_CZGRIDX": _overlap(new[13], _long(candidates, column="bbg__building_permits_raw_candidate__as_reported")),
        "25_vs_previous_local_trade_balance": _overlap(new[25], _long(candidates, column="local__trade_balance__as_stored")),
        "26_vs_czso_CEN0303_month_on_month_pct": _overlap(new[26] - 100.0, _long(exact, a6_number=26)),
        "43_vs_previous_czso_yoy": _overlap(new[43], _long(exact, a6_number=43)),
        "52_backcast_test_anchor_2015": backcast_check(mm52, yoy_total),
        "54_vs_bloomberg_month_end": _overlap(new[54], _long(candidates, column="bbg__pribor_3m__month_last")),
        "58_vs_previous_exact": _overlap(new[58], _long(exact, a6_number=58)),
        "59_vs_previous_exact": _overlap(new[59], _long(exact, a6_number=59)),
        "60_vs_bloomberg_CO1_month_mean": _overlap(new[60], _long(candidates, column="bbg__brent_front_reference__month_mean")),
        "61_vs_bloomberg_TTF_day_ahead": _overlap(new[61], _long(candidates, column="bbg__gas_day_ahead_spot__month_mean"), log_changes=True),
        "62_vs_bloomberg_industrial_metals": _overlap(new[62], _long(candidates, column="bbg__industrial_metals_spot_index__month_mean"), log_changes=True),
        "63_vs_bloomberg_agriculture": _overlap(new[63], _long(candidates, column="bbg__agriculture_spot_index__month_mean"), log_changes=True),
    }
    manifest = {
        "tool": "build_exact_overrides", "created_at": datetime.now(timezone.utc).isoformat(),
        "downloads": downloads | {"czso_permits": {"file": str(PERMITS.relative_to(ROOT)), "sha256": sha256_file(PERMITS)}},
        "arad_snapshot_ids": snapshots, "observed_to": {"monthly_report_database": str(LAST_MONTH), "quarterly_report_database": str(LAST_QUARTER)},
        "rows": {str(k): v for k, v in sorted(sources.items())}, "seasonal_policies": {str(k): v for k, v in POLICIES.items()},
        "checks": checks, "output": {"file": "overrides_long.csv", "sha256": sha256_file(output / "overrides_long.csv"), "rows": int(len(frame))},
        "tool_sha256": sha256_file(Path(__file__)),
    }
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2, default=str), encoding="utf-8")
    return manifest


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--replace", action="store_true", help="rewrite an existing output folder")
    args = parser.parse_args(argv)
    manifest = build(args.output, replace=args.replace)
    print(json.dumps({"rows": manifest["rows"], "checks": manifest["checks"]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
