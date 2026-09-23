"""Offline parsing and provenance checks for the official A6 comparison sources."""
from datetime import datetime
import io
import json
import zipfile

import openpyxl
import pytest

from tools.paper_replication import official_sources as src


def _payload():
    return {"id": ["geo", "time"], "size": [2, 3],
            "dimension": {"geo": {"category": {"index": {"CZ": 0, "DE": 1}}},
                          "time": {"category": {"index": {"2026-01": 0, "2026-02": 1, "2026-03": 2}}}},
            "value": {"0": 1.5, "2": 2.5, "4": -3.0}, "status": {"4": "p"}}


def test_jsonstat_keeps_only_published_cells_and_status_flags():
    frame = src.jsonstat_to_frame(_payload())
    assert frame[["geo", "time", "value"]].values.tolist() == [
        ["CZ", "2026-01", 1.5], ["CZ", "2026-03", 2.5], ["DE", "2026-02", -3.0]]
    assert frame["status"].isna().tolist() == [True, True, False]
    assert frame["status"].iloc[2] == "p"


def _archive(folder, name="services_total_sa_nace2.zip"):
    wb = openpyxl.Workbook()
    index = wb.active
    index.title = "Index"
    index.append([None, "Series name"])
    index.append(["Monthly Questions"])
    index.append(["1", "Business situation development over the past 3 months"])
    monthly = wb.create_sheet("SERVICES MONTHLY")
    monthly.append([None, None, "SERV.CZ.TOT.1.BS.M", "SERV.DE.TOT.1.BS.M", None, "SERV.CZ.TOT.2.B.M"])
    monthly.append([datetime(2026, 1, 31), None, 10.5, "NA", None, 3.0])
    monthly.append([datetime(2026, 2, 28), None, 11.0, 4.0, None, "NA"])
    quarterly = wb.create_sheet("SERVICES QUARTERLY")
    quarterly.append([None, None, "SERV.CZ.TOT.7.F2S.Q"])
    quarterly.append(["2026-Q1", None, 40.0])
    buffer = io.BytesIO()
    wb.save(buffer)
    path = folder / name
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(name.replace(".zip", ".xlsx"), buffer.getvalue())
    return path


def test_ecfin_archive_is_long_with_adjustment_flag_and_no_filled_cells(tmp_path):
    out = src.read_ecfin_archive(_archive(tmp_path))
    assert set(zip(out.series_code, out.period)) == {
        ("SERV.CZ.TOT.1.BS.M", "2026-01"), ("SERV.CZ.TOT.1.BS.M", "2026-02"),
        ("SERV.DE.TOT.1.BS.M", "2026-02"), ("SERV.CZ.TOT.2.B.M", "2026-01"),
        ("SERV.CZ.TOT.7.F2S.Q", "2026-Q1")}
    adjusted = dict(zip(out.series_code, out.seasonally_adjusted))
    assert adjusted["SERV.CZ.TOT.1.BS.M"] and adjusted["SERV.CZ.TOT.7.F2S.Q"]
    assert not adjusted["SERV.CZ.TOT.2.B.M"]
    units = dict(zip(out.series_code, out.unit))
    assert units["SERV.CZ.TOT.1.BS.M"] == "balance" and units["SERV.CZ.TOT.7.F2S.Q"] == "percent_of_firms"


def test_ecfin_geography_filter(tmp_path):
    out = src.read_ecfin_archive(_archive(tmp_path), geos=["DE"])
    assert out.series_code.unique().tolist() == ["SERV.DE.TOT.1.BS.M"]


@pytest.mark.parametrize("newline", ["\r\n", "\n", "\r\r\n"])
def test_headers_parser_uses_the_final_response(newline):
    text = newline.join(["HTTP/1.1 301 Moved", "Location: x", "", "HTTP/1.1 200 OK", 'ETag: "e"',
                         "Last-Modified: Tue, 25 Aug 2026 16:47:08 GMT", "", ""])
    headers = src._parse_headers(text)
    assert headers["status_line"] == "HTTP/1.1 200 OK"
    assert headers["etag"] == '"e"' and headers["last-modified"] == "Tue, 25 Aug 2026 16:47:08 GMT"


def test_ecfin_manifest_refuses_to_describe_a_changed_archive(tmp_path):
    folder = tmp_path / "nace2_ecfin_2608"
    folder.mkdir()
    path = _archive(folder)
    (folder / f"{path.name}.headers.txt").write_bytes(
        b"HTTP/1.1 200 OK\r\nLast-Modified: Tue, 25 Aug 2026 16:47:08 GMT\r\n\r\n")
    manifest = src.ecfin_manifest(folder)
    assert manifest["archives"][path.name]["last_modified"] == "Tue, 25 Aug 2026 16:47:08 GMT"
    assert manifest["archives"][path.name]["url"].endswith("nace2_ecfin_2608/services_total_sa_nace2.zip")
    assert src.ecfin_manifest(folder) == manifest  # unchanged archives: identical manifest
    path.write_bytes(path.read_bytes() + b"x")
    with pytest.raises(ValueError, match="immutable"):
        src.ecfin_manifest(folder)


class _Response:
    def __init__(self, payload):
        self._payload = payload
        self.content = json.dumps(payload).encode()
        self.url = "https://example.invalid/query"

    def raise_for_status(self):
        return None

    def json(self):
        return self._payload


class _Session:
    def __init__(self, payloads):
        self.payloads = payloads

    def get(self, url, params=None, timeout=None):
        return _Response(self.payloads[url.rsplit("/", 1)[-1]])


def test_eurostat_download_records_an_empty_query_as_an_error(tmp_path):
    empty = {"id": ["geo"], "size": [1], "dimension": {"geo": {"category": {"index": {"CZ": 0}}}}, "value": {}}
    queries = {"good": ("ds_good", {}), "empty": ("ds_empty", {})}
    manifest = src.download_eurostat(tmp_path / "out", queries, session=_Session({"ds_good": _payload(), "ds_empty": empty}))
    assert manifest["status"] == "partial"
    assert manifest["queries"]["good"]["rows"] == 3
    assert [e["query"] for e in manifest["errors"]] == ["empty"]
    with pytest.raises(FileExistsError):
        src.download_eurostat(tmp_path / "out", queries, session=_Session({}))
